#!/usr/bin/env python3
"""Generuje szkic decku HELVIA: DOKŁADNIE 3 slajdy (E1 → slajd 1, E2 → slajd 2, E3 + E4 → slajd 3).

Wszystko to natywne kształty i pola tekstowe (edycja w PowerPoint / Keynote / import do Canva).
Pola oznaczone [DO DECYZJI] i [E4] to miejsca, gdzie zespół jeszcze nie zdecydował.

Użycie:  python3 narzedzia/generuj_deck.py [wyjście.pptx]
"""
import sys
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

OUT = Path(sys.argv[1]) if len(sys.argv) > 1 else \
    Path(__file__).resolve().parent.parent / "05_etap5_prezentacja" / "slajdy" / "HELVIA_deck_v1.pptx"

NAVY = RGBColor(0x1F, 0x3A, 0x5F)
TEAL = RGBColor(0x0F, 0x7C, 0x7E)
RED = RGBColor(0xB8, 0x3B, 0x2E)
AMBER = RGBColor(0xC7, 0x7C, 0x0E)
INK = RGBColor(0x1C, 0x24, 0x30)
MUTED = RGBColor(0x5B, 0x67, 0x75)
LINE = RGBColor(0xD5, 0xDB, 0xE2)
PANEL = RGBColor(0xF3, 0xF5, 0xF8)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
AMBER_BG = RGBColor(0xFD, 0xF3, 0xE1)

prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
BLANK = prs.slide_layouts[6]


def box(s, x, y, w, h, fill=None, line=None, shape=MSO_SHAPE.RECTANGLE, lw=0.75):
    sh = s.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    if fill is None:
        sh.fill.background()
    else:
        sh.fill.solid()
        sh.fill.fore_color.rgb = fill
    if line is None:
        sh.line.fill.background()
    else:
        sh.line.color.rgb = line
        sh.line.width = Pt(lw)
    sh.shadow.inherit = False
    return sh


def text(s, x, y, w, h, runs, size=12, color=INK, bold=False, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, shape=None):
    """runs: str | list[paragraph]; paragraph: str | list[(tekst, {opcje})]"""
    tb = shape if shape is not None else s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    for m in ("margin_left", "margin_right"):
        setattr(tf, m, Inches(0.08))
    tf.margin_top = tf.margin_bottom = Inches(0.04)
    paras = runs if isinstance(runs, list) else [runs]
    for i, p in enumerate(paras):
        para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        para.alignment = align
        segs = p if isinstance(p, list) else [(p, {})]
        for seg, o in segs:
            r = para.add_run()
            r.text = seg
            f = r.font
            f.name = "Calibri"
            f.size = Pt(o.get("size", size))
            f.bold = o.get("bold", bold)
            f.italic = o.get("italic", False)
            f.color.rgb = o.get("color", color)
        if isinstance(p, list) and p and "space" in p[0][1]:
            para.space_after = Pt(p[0][1]["space"])
    return tb


def header(s, n, etap, headline):
    box(s, 0, 0, 13.333, 0.09, fill=NAVY)
    text(s, 0.55, 0.28, 9, 0.3, [[(f"{n}/3  ·  {etap}", {"bold": True, "color": TEAL, "size": 12})]])
    text(s, 0.55, 0.58, 12.2, 1.1, headline, size=26, bold=True, color=NAVY)


def footer(s, src, status=None):
    box(s, 0.55, 7.0, 12.23, 0.01, fill=LINE)
    text(s, 0.55, 7.04, 9.6, 0.35, src, size=9, color=MUTED)
    if status:
        tag(s, 10.45, 7.06, 2.33, status)


def tag(s, x, y, w, label, h=0.28):
    b = box(s, x, y, w, h, fill=AMBER_BG, line=AMBER, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    text(None, 0, 0, 0, 0, [[(label, {"bold": True, "color": AMBER, "size": 9})]],
         align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, shape=b)


def kpi(s, x, y, w, big, cap, src, col=NAVY, h=1.42):
    box(s, x, y, w, h, fill=PANEL)
    box(s, x, y, 0.07, h, fill=col)
    text(s, x + 0.18, y + 0.08, w - 0.25, 0.6, [[(big, {"bold": True, "size": 30, "color": col})]])
    text(s, x + 0.18, y + 0.68, w - 0.25, 0.5, cap, size=12, color=INK)
    text(s, x + 0.18, y + h - 0.32, w - 0.25, 0.28, src, size=8.5, color=MUTED)


def notes(s, t):
    s.notes_slide.notes_text_frame.text = t


# ───────────────────────────── SLAJD 1 — DIAGNOZA (E1 → 1C)
s = prs.slides.add_slide(BLANK)
header(s, 1, "DIAGNOZA · JAK JEST?",
       "HELVIA traci 2/3 zysku operacyjnego na karach SLA — problemem jest jakość obsługi, nie cena")
kpi(s, 0.55, 1.85, 3.9, "6,27 mln zł", "kar SLA w 12 mies. = 65% EBIT (9,64 mln zł)", "Klienci!O2:O15 · Dashboard!A8", RED)
kpi(s, 0.55, 3.42, 3.9, "94%", "kar płacą 2 klienci: NORDA (4,93 mln) i Lumea (0,99 mln)", "Klienci!O2, O9", RED)
kpi(s, 0.55, 4.99, 3.9, "24,5%", "rotacji w Obsłudze Klienta (2,7× mediany) · SLA 92,1% · 53% błędów firmy",
    "Działy!E19, D41, C41", NAVY)

# łańcuch przyczyn
text(s, 4.85, 1.85, 4.3, 0.3, [[("Łańcuch przyczyn", {"bold": True, "size": 13, "color": NAVY})]])
chain = [("Kary SLA 6,27 mln zł", "marża NORDY 4,6% zamiast 23% przed karami", RED),
         ("SLA poniżej umów", "NORDA 91,0% (szczyt XI 2025: 88,9%) · Lumea 92,4%", RED),
         ("Przeciążona Obsługa Klienta", "34% juniorów · 23% ze stażem < 1 roku · eNPS −64", NAVY),
         ("Lęk i brak wsparcia", "82% boi się o etat (Q20) · przeszkolonych z AI: 9%", NAVY)]
y = 2.22
for i, (t1, t2, c) in enumerate(chain):
    b = box(s, 4.85, y, 4.3, 0.88, fill=WHITE, line=LINE)
    box(s, 4.85, y, 0.07, 0.88, fill=c)
    text(s, 5.0, y + 0.06, 4.1, 0.8, [[(t1, {"bold": True, "size": 13, "color": c})], [(t2, {"size": 11})]])
    if i < 3:
        a = s.shapes.add_shape(MSO_SHAPE.UP_ARROW, Inches(6.88), Inches(y + 0.9), Inches(0.24), Inches(0.22))
        a.fill.solid(); a.fill.fore_color.rgb = LINE; a.line.fill.background()
    y += 1.12
text(s, 4.85, 6.6, 4.3, 0.3, "Strzałka: co jest przyczyną czego (od dołu do góry)", size=9, color=MUTED)

# priorytety
box(s, 9.5, 1.85, 3.28, 4.95, fill=PANEL)
text(s, 9.65, 1.95, 3.0, 0.3, [[("3 priorytety zarządu", {"bold": True, "size": 13, "color": NAVY})]])
pri = [("1", "Jakość SLA u NORDY i Lumei", "przed rozmową o cenie"),
       ("2", "Ludzie w Obsłudze Klienta", "stabilny zespół = mniej błędów"),
       ("3", "Odciążenie Finansów", "140 h nadgodzin/mies. · obciążenie Q09 = 2,6")]
y = 2.42
for n, t1, t2 in pri:
    c = box(s, 9.68, y, 0.42, 0.42, fill=TEAL, shape=MSO_SHAPE.OVAL)
    text(None, 0, 0, 0, 0, [[(n, {"bold": True, "color": WHITE, "size": 14})]], align=PP_ALIGN.CENTER,
         anchor=MSO_ANCHOR.MIDDLE, shape=c)
    text(s, 10.2, y - 0.04, 2.5, 1.0, [[(t1, {"bold": True, "size": 13})], [(t2, {"size": 10.5, "color": MUTED})]])
    y += 1.05
text(s, 9.65, 5.6, 3.0, 1.1,
     [[("Świadomie odkładamy: ", {"bold": True, "size": 10.5}), ("ogon nierentownych klientów (Ibera −1,8%, Sørvik 2,1%) — wracamy przy odnowieniach w 2027.", {"size": 10.5})]])
tag(s, 9.65, 1.5, 3.0, "1C DO POTWIERDZENIA PRZEZ ZESPÓŁ")
footer(s, "Źródła: HITL_Etap_1 — Klienci, Działy, Historia miesięczna, ankieta (513 odp.). Liczby z 1B (S1, S3, S4–S5, S8).")
notes(s, """[0:00–0:20 · Strateg] Problem HELVII to nie technologia ani cena, tylko jakość obsługi. Kary SLA zabrały w rok 6,27 mln zł — dwie trzecie zysku operacyjnego.

[0:20–1:00 · Analityk danych] 94% tych kar płacą dwaj klienci: NORDA i Lumea. Obsługuje ich głównie Obsługa Klienta — dział z rotacją 24,5%, najniższym SLA w firmie i ponad połową wszystkich błędów. Co trzecia osoba to junior, prawie co czwarta pracuje krócej niż rok, a 82% boi się, że AI zabierze im pracę. Przeszkolonych z AI jest tam 9%. Dlatego nasze trzy priorytety to: jakość u NORDY i Lumei, ludzie w Obsłudze Klienta i odciążenie Finansów, które mają najwięcej nadgodzin w firmie.

[DO DECYZJI] Arkusz 1C w pliku jest pusty — potwierdźcie 3 priorytety i wpiszcie je w 1C (S1, S3, S8 z waszego 1B).""")

# ───────────────────────────── SLAJD 2 — DECYZJA (E2 → 2B)
s = prs.slides.add_slide(BLANK)
header(s, 2, "DECYZJA · JAK MOŻE BYĆ?",
       "Nie kupujemy lojalności NORDY rabatem — naprawiamy jakość i rośniemy tam, gdzie etat zarabia najwięcej")

# lewa kolumna: NORDA
text(s, 0.55, 1.85, 5.9, 0.3, [[("NORDA żąda −12% ceny. Odpowiedź: negocjujemy", {"bold": True, "size": 14, "color": NAVY})]])
rows = [("Rabat", "Stawka / op.", "Marża NORDY / rok"),
        ("0% teraz + 90-dniowy plan SLA", "38,70 zł", "+1,22 mln zł"),
        ("4% od II kw. 2027 (za warunki)", "37,15 zł", "+0,16 mln zł"),
        ("4,6% — próg rentowności", "36,92 zł", "0 zł"),
        ("12% — żądanie NORDY", "34,06 zł", "−1,96 mln zł")]
cols = [(0.55, 2.9), (3.45, 1.25), (4.7, 1.75)]
y = 2.25
for i, r in enumerate(rows):
    hdr = i == 0
    bad = i == 4
    fill = NAVY if hdr else (RGBColor(0xFB, 0xEA, 0xE7) if bad else (PANEL if i % 2 else WHITE))
    box(s, 0.55, y, 5.9, 0.42, fill=fill)
    for j, (cx, cw) in enumerate(cols):
        col = WHITE if hdr else (RED if bad and j == 2 else INK)
        text(s, cx, y + 0.04, cw, 0.36, [[(r[j], {"bold": hdr or j == 2, "size": 11.5, "color": col})]],
             align=PP_ALIGN.LEFT if j == 0 else PP_ALIGN.RIGHT)
    y += 0.44
text(s, 0.55, 4.5, 5.9, 1.4, [
    [("W zamian oczekujemy: ", {"bold": True, "size": 11.5}),
     ("przedłużenia o 3 lata, minimalnego wolumenu, prognozy szczytów, limitu kar i podziału oszczędności z automatyzacji.", {"size": 11.5})],
    [("Argument: ", {"bold": True, "size": 11.5}),
     ("kary SLA NORDY (4,93 mln zł) są 4× większe niż jej marża — naprawa jakości daje więcej niż jakikolwiek ruch cenowy.", {"size": 11.5})]])
box(s, 0.55, 5.95, 5.9, 0.85, fill=PANEL)
text(s, 0.7, 6.0, 5.65, 0.8, [[("Plan B: ", {"bold": True, "size": 11, "color": RED}),
                               ("nie schodzimy poniżej 4,6%. Jeśli NORDA odejdzie, ok. 51 etatów przechodzi do S-02, S-01 i S-10 (ok. 15) "
                                "oraz naturalnej rotacji — przesunięcia zamiast zwolnień.", {"size": 11})]])

# prawa kolumna: kierunek
box(s, 6.85, 1.85, 5.93, 4.95, fill=PANEL)
text(s, 7.0, 1.95, 5.7, 0.6, [[("Kierunek: wartość zamiast wolumenu", {"bold": True, "size": 14, "color": NAVY})],
                              [("analityka i usługi oparte na AI dla obecnych klientów, realizowane przez działy gotowe na zmianę", {"size": 11, "color": MUTED})]])
proj = [("S-03", "Analityka operacyjna", "Quenta", "80% · marża 21% · 2,1 etatu"),
        ("S-20", "Analiza danych", "Quenta", "70% · marża 23% · 3,0 etatu"),
        ("S-14", "Lejek sprzedaży i ofertowanie", "Pilawa", "80% · marża 18% · 3,5 etatu")]
y = 2.68
for pid, n, k, d in proj:
    box(s, 7.0, y, 5.63, 0.5, fill=WHITE, line=LINE)
    text(s, 7.08, y + 0.06, 0.7, 0.4, [[(pid, {"bold": True, "size": 12, "color": TEAL})]])
    text(s, 7.75, y + 0.06, 2.9, 0.4, [[(n, {"bold": True, "size": 11.5}), (f" · {k}", {"size": 11.5, "color": MUTED})]])
    text(s, 10.3, y + 0.06, 2.3, 0.4, d, size=10.5, color=MUTED, align=PP_ALIGN.RIGHT)
    y += 0.56
# słupki marża na etat
text(s, 7.0, 4.42, 5.6, 0.3, [[("Oczekiwana marża na etat (tys. zł)", {"bold": True, "size": 11.5})]])
bars = [("Nasz kierunek", 87, TEAL), ("Projekty Finansów", 41, MUTED), ("Projekty Obsługi Klienta", 17, RED)]
y = 4.78
for lab, v, c in bars:
    text(s, 7.0, y, 1.95, 0.32, lab, size=10.5)
    box(s, 9.0, y + 0.05, 3.1 * v / 87, 0.24, fill=c)
    text(s, 9.05 + 3.1 * v / 87, y, 0.6, 0.32, [[(str(v), {"bold": True, "size": 11, "color": c})]])
    y += 0.38
text(s, 7.0, 5.95, 5.7, 0.85, [[("Kosztem: ", {"bold": True, "size": 11, "color": RED}),
                                ("ekspansji Obsługi Klienta (S-02, S-09, S-12, S-19: ok. 25 etatów), nowych dużych klientów (S-06, S-22) "
                                 "i przesunięcia S-01 Kastella na II kw. 2027 — po pilotażu automatyzacji w Finansach.", {"size": 11})]])
footer(s, "Źródła: FINAL_HIL_Etap_2 — 2A G37:L43 (scenariusze NORDY), 2A A6:L27 (pipeline: wartość × szansa × marża / etaty), 2B C12–C22.",
       "RYZYKO: 3,7 MLN vs 26,4 MLN NORDY")
notes(s, """[1:00–1:50 · Strateg] NORDA żąda obniżki o 12%. Przy marży 4,6% każdy rabat powyżej 4,6% oznacza stratę na każdej operacji — 12% to minus 1,96 mln zł rocznie. Dlatego negocjujemy: teraz zero rabatu i 90-dniowy plan naprawy SLA, a od drugiego kwartału 2027 do 4% — w zamian za przedłużenie o 3 lata, prognozę szczytów i limit kar. Kary NORDY są cztery razy większe niż jej marża, więc jakość daje więcej niż cena.

Jednocześnie stawiamy na wartość zamiast wolumenu: trzy projekty analityczne i sprzedażowe dla Quenty i Pilawy. Dają 87 tys. zł marży na etat, pięć razy więcej niż projekty w Obsłudze Klienta. Robimy to kosztem rozbudowy Obsługi Klienta i nowych dużych klientów.

[DO WERYFIKACJI] Start S-03 w danych to 24.12.2027 przy decyzji do 15.11.2026 — potwierdźcie z mentorem. Komórka 2B C24 pokazuje #NAME? (funkcja MINIFS) — sprawdźcie w Excelu przed oddaniem.""")

# ───────────────────────────── SLAJD 3 — CZŁOWIEK + AI + LUDZIE (E3 + E4)
s = prs.slides.add_slide(BLANK)
header(s, 3, "CZŁOWIEK + AI · A CO Z LUDŹMI?",
       "AI przejmuje rutynę w 3 procesach, człowiek zostaje przy decyzjach — uwolniony czas idzie w jakość i wzrost, nie w zwolnienia")

# tabela procesów
text(s, 0.55, 1.85, 7.0, 0.3, [[("Pierwsza fala AI (Etap 3)", {"bold": True, "size": 14, "color": NAVY})]])
tag(s, 3.7, 1.86, 3.85, "HITL: P-012 ✓  ·  P-001, P-039, P-015 DO DECYZJI")
prow = [("Proces", "Tryb", "Uwolnione h/m*", "Po co (strategia)"),
        ("P-001 Rejestracja faktur", "Automat + wyjątki", "3 554 (22 et.)", "Odciąża Finanse → odblokowuje S-01"),
        ("P-012 E-maile klientów", "HITL", "2 825 (18 et.)", "Plan SLA dla NORDY w szczytach"),
        ("P-039 Oferty handlowe", "Asystent", "854 (5 et.)", "Wspiera S-14 Pilawa"),
        ("P-015 Reklamacje złożone", "Zostaje u człowieka", "—", "Empatia i decyzja; 31% w OK: „trudne rozmowy”")]
cw = [(0.55, 2.25), (2.8, 1.45), (4.25, 1.3), (5.55, 2.0)]
y = 2.25
for i, r in enumerate(prow):
    hdr = i == 0
    human = i == 4
    fill = NAVY if hdr else (RGBColor(0xE6, 0xF2, 0xF2) if human else (PANEL if i % 2 else WHITE))
    box(s, 0.55, y, 7.0, 0.5, fill=fill)
    for j, (cx, w) in enumerate(cw):
        col = WHITE if hdr else (TEAL if human and j < 2 else INK)
        text(s, cx, y + 0.03, w, 0.46, [[(r[j], {"bold": hdr or j == 0, "size": 10.5 if j == 3 else 11, "color": col})]],
             anchor=MSO_ANCHOR.MIDDLE)
    y += 0.52

# mini HITL
text(s, 0.55, 4.95, 7.0, 0.3, [[("HITL dla P-012 (CRM Helios): AI przygotowuje, człowiek decyduje", {"bold": True, "size": 12, "color": NAVY})]])
steps = [("1 · AI", "rozpoznaje temat, klienta i zamówienie; dzieli e-mail z kilkoma pytaniami", TEAL),
         ("2 · AI", "pisze szkic z historii kontaktów i bazy gotowych odpowiedzi", TEAL),
         ("3 · Konsultant", "zatwierdza lub poprawia — wysyłka zawsze po decyzji człowieka", NAVY),
         ("Eskalacja → lider", "skarga, reklamacja, roszczenie, groźba prawna, dane osobowe, prośba o wyjątek", AMBER)]
x = 0.55
for i, (t1, t2, c) in enumerate(steps):
    box(s, x, 5.3, 1.6, 1.0, fill=WHITE, line=c, lw=1.25)
    text(s, x + 0.04, 5.33, 1.54, 0.96, [[(t1, {"bold": True, "size": 11, "color": c})], [(t2, {"size": 8.5})]])
    if i < 3:
        a = s.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(x + 1.61), Inches(5.69), Inches(0.17), Inches(0.2))
        a.fill.solid(); a.fill.fore_color.rgb = LINE; a.line.fill.background()
    x += 1.8
text(s, 0.55, 6.32, 7.0, 0.62, [[("Koszt technologii rok 1: ", {"bold": True, "size": 11}),
                                 ("ok. 295 tys. zł z limitu 300 tys. zł ", {"size": 11}),
                                 ("(3D: Operacje 111 tys. · Hybryda 150 tys. · Asystent 33 tys.)", {"size": 10, "color": MUTED})]])

# prawa: A co z ludźmi (E4)
box(s, 7.85, 1.85, 4.93, 4.95, fill=AMBER_BG)
text(s, 8.0, 1.95, 4.7, 0.3, [[("A co z ludźmi? (Etap 4)", {"bold": True, "size": 14, "color": NAVY})]])
tag(s, 10.95, 1.97, 1.7, "UZUPEŁNIĆ W E4")
text(s, 8.0, 2.4, 4.7, 0.9, [[("ok. 45 etatów-ekw.", {"bold": True, "size": 26, "color": TEAL})],
                             [("uwolnionego czasu miesięcznie (hipoteza*)", {"size": 11, "color": MUTED})]])
plan = [("Gdzie trafia czas", "stabilizacja SLA w szczytach NORDY · S-01 Kastella (6,5 et.) · projekty kierunku (8,6 et.)"),
        ("Zero zwolnień", "dopasowanie przez naturalną rotację i wstrzymanie rekrutacji"),
        ("Szkolenia", "Obsługa Klienta: 9% → [cel] · Finanse: 17% → [cel] · IT jako ambasadorzy (73%)"),
        ("Komunikacja", "polityka AI i jasne role — dziś 6,8% widzi jasną komunikację (Q21)"),
        ("3 KPI na 12 mies.", "SLA OK 92,1% → [cel] · kary −[x]% · rotacja OK 24,5% → [cel]")]
y = 3.42
for t1, t2 in plan:
    text(s, 8.0, y, 4.7, 0.62, [[(t1 + ": ", {"bold": True, "size": 11}), (t2, {"size": 10.5})]])
    y += 0.66
footer(s, "* Hipoteza skrócenia czasu [ZAŁOŻENIE]: P-001 −60%, P-012 −40%, P-039 −30% godzin z 3A; 1 etat = 160 h/m. "
          "Źródła: Etap 3 — 3A (wolumeny, godziny), 3D (cennik), karty procesów DOCX; ankieta Q16, Q21, Q26.")
notes(s, """[1:50–2:40 · Analityk ludzi i ryzyk] W pierwszej fali AI bierze rutynę w trzech procesach. Rejestracja faktur działa automatycznie z wyjątkami dla człowieka — to odciąża Finanse i odblokowuje projekt dla Kastelli. Odpowiedzi na e-maile klientów projektujemy w trybie human in the loop: AI rozpoznaje sprawę i pisze szkic odpowiedzi, ale wysyłka zawsze wymaga decyzji konsultanta, a skargi, reklamacje, roszczenia i prośby o wyjątek od razu trafiają do lidera zespołu. Reklamacje złożone świadomie zostawiamy ludziom — tam liczą się empatia i decyzja. Koszt technologii: ok. 295 tys. zł w pierwszym roku.

Uwolniony czas — ok. 45 etatów — nie oznacza zwolnień. Idzie na jakość w szczytach NORDY i na nowe projekty, a zatrudnienie dopasowuje naturalna rotacja. Szkolimy Obsługę Klienta, w której dziś przeszkolonych jest 9%.

[2:40–3:00 · Decydent] Zarząd odzyskuje marżę, pracownicy dostają rozwój i bezpieczeństwo, klienci — lepszą obsługę. To nie jest AI wszędzie. To AI tam, gdzie się zwraca, i człowiek tam, gdzie decyduje. Dziękujemy.

[DECYZJA ZESPOŁU] Proces HITL w 3C: P-012 (e-maile klientów). [DO DECYZJI] Pozostałe procesy pierwszej fali (P-001, P-039), proces dla człowieka (P-015) i % skrócenia czasu to propozycja asystenta. Prawa kolumna to szkielet na Etap 4.""")

OUT.parent.mkdir(parents=True, exist_ok=True)
prs.save(OUT)
assert len(prs.slides) == 3, "Deck musi mieć dokładnie 3 slajdy"
print(f"Zapisano: {OUT} ({len(prs.slides)} slajdy)")
