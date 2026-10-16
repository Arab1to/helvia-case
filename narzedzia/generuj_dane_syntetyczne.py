#!/usr/bin/env python3
"""Fikcyjna HELVIA do próby generalnej (8 działów, 14 klientów, pipeline, ankieta, karty procesów).

⚠️ DANE SYNTETYCZNE — wymyślone, tylko do treningu tempa E1–E2. Nigdy na slajd.
Sumy kontrolne zgodne z opisem case'u: Σ FTE = 720, Σ przychodów = 114,93 mln zł.

Użycie:  python3 narzedzia/generuj_dane_syntetyczne.py [--seed 7] [ścieżka.xlsx]
"""
import argparse
import random
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

ROOT = Path(__file__).resolve().parent.parent
DZIALY = ["Finanse i księgowość", "HR i payroll", "Obsługa klienta", "IT", "Logistyka", "Administracja",
          "Zakupy i rozliczenia", "Analityka i raportowanie"]
PROCESY = [("Obsługa faktur kosztowych", "Finanse i księgowość"), ("Uzgadnianie sald", "Finanse i księgowość"),
           ("Naliczanie płac", "HR i payroll"), ("Rekrutacja — preselekcja", "HR i payroll"),
           ("Odpowiedzi na zapytania e-mail", "Obsługa klienta"), ("Reklamacje i eskalacje", "Obsługa klienta"),
           ("Helpdesk L1", "IT"), ("Awizacje i dokumenty przewozowe", "Logistyka"),
           ("Raportowanie miesięczne dla klientów", "Analityka i raportowanie"), ("Obieg dokumentów", "Administracja")]


def podziel(suma, n, rng, lo, hi, nd):
    w = [rng.uniform(lo, hi) for _ in range(n)]
    s = sum(w)
    vals = [round(x / s * suma, nd) for x in w]
    vals[-1] = round(suma - sum(vals[:-1]), nd)
    return vals


def main():
    p = argparse.ArgumentParser()
    p.add_argument("--seed", type=int, default=7)
    p.add_argument("out", nargs="?", default=str(ROOT / "dane" / "syntetyczne" / "HELVIA_syntetyczne.xlsx"))
    a = p.parse_args()
    rng = random.Random(a.seed)
    wb = Workbook()
    warn = Font(bold=True, color="C00000")

    def ark(title, headers):
        ws = wb.create_sheet(title)
        ws.append(["DANE SYNTETYCZNE — FIKCYJNE, tylko do próby generalnej"])
        ws["A1"].font = warn
        ws.append(headers)
        for c in ws[2]:
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = PatternFill("solid", fgColor="1F3A5F")
        return ws

    wb.remove(wb.active)
    # Działy — jeden dział „gorący” (wysoka rotacja + SLA poniżej umowy), żeby H1 dało się przetestować
    goracy = rng.choice(["Obsługa klienta", "HR i payroll", "Logistyka"])
    fte = podziel(720, 8, rng, 0.6, 1.6, 0)
    ws = ark("Dzialy", ["Dział", "FTE", "Obciążenie (h/FTE/mies.)", "Wydajność (sprawy/FTE/dzień)", "SLA osiągnięte %",
                        "SLA w umowie %", "Rotacja roczna %", "Nadgodziny h/mies.", "Koszt etatu mies. (zł)"])
    for d, f in zip(DZIALY, fte):
        hot = d == goracy
        sla_umowa = rng.choice([95, 97, 98])
        ws.append([d, int(f), rng.randint(165, 185) if hot else rng.randint(150, 172), round(rng.uniform(18, 42), 1),
                   round(sla_umowa - (rng.uniform(3.5, 6) if hot else rng.uniform(-1.5, 2.5)), 1), sla_umowa,
                   round(rng.uniform(0.26, 0.34) if hot else rng.uniform(0.09, 0.19), 3),
                   rng.randint(12, 22) if hot else rng.randint(2, 9), rng.randint(8500, 13500)])

    # Klienci
    przych = sorted(podziel(114.93, 14, rng, 0.25, 2.2, 2), reverse=True)
    ws = ark("Klienci", ["Klient", "Przychód LTM (mln zł)", "Marża klienta %", "SLA osiągnięte %", "SLA w umowie %",
                         "Koniec kontraktu (mies.)", "Presja cenowa (1–5)", "Główny dział"])
    for i, pr in enumerate(przych, 1):
        dz = goracy if i <= 3 and rng.random() < 0.7 else rng.choice(DZIALY)
        umowa = rng.choice([95, 97, 98])
        ws.append([f"Klient {chr(64 + i)}", pr, round(rng.uniform(-0.02, 0.15), 3),
                   round(umowa - (rng.uniform(2, 5) if dz == goracy else rng.uniform(-1.5, 2)), 1), umowa,
                   rng.randint(3, 30), rng.randint(1, 5), dz])

    # Pipeline
    ws = ark("Pipeline", ["Kontrakt", "Wartość roczna (mln zł)", "Prawdopodobieństwo %", "FTE potrzebne", "Dział realizujący", "Start (mies.)"])
    for i in range(1, 7):
        ws.append([f"P{i}", round(rng.uniform(1.5, 7.5), 2), rng.choice([0.3, 0.4, 0.5, 0.6, 0.7, 0.8]),
                   rng.randint(8, 30), rng.choice(DZIALY), rng.randint(2, 9)])

    # Ankieta — tylko agregaty per dział (tak jak wklejamy do Claude)
    ws = ark("Ankieta_agregaty", ["Dział", "Liczba odpowiedzi", "Obawa przed AI: % ocen 4–5", "Gotowość do szkoleń: % ocen 4–5", "eNPS"])
    for d, f in zip(DZIALY, fte):
        hot = d == goracy
        ws.append([d, int(f * rng.uniform(0.55, 0.8)), round(rng.uniform(0.48, 0.66) if hot else rng.uniform(0.2, 0.42), 2),
                   round(rng.uniform(0.4, 0.75), 2), rng.randint(-35, -10) if hot else rng.randint(-10, 25)])

    # Karty procesów
    ws = ark("Procesy", ["Proces", "Dział", "Wolumen / mies.", "Czas na sprawę (min)", "Udział kroków powtarzalnych %",
                         "Dane ustrukturyzowane %", "Błędy %", "Dane osobowe (T/N)", "Decyzja wobec ludzi (T/N)"])
    for pr, dz in PROCESY:
        osob = "T" if dz in ("HR i payroll", "Obsługa klienta") else rng.choice("TN")
        ws.append([pr, dz, rng.randint(400, 9000), rng.randint(4, 35), round(rng.uniform(0.4, 0.9), 2),
                   round(rng.uniform(0.3, 0.95), 2), round(rng.uniform(0.01, 0.08), 3), osob,
                   "T" if pr in ("Naliczanie płac", "Rekrutacja — preselekcja", "Reklamacje i eskalacje") else "N"])

    for s in wb.worksheets:
        for col in "ABCDEFGHI":
            s.column_dimensions[col].width = 22
    Path(a.out).parent.mkdir(parents=True, exist_ok=True)
    wb.save(a.out)
    print(f"Zapisano: {a.out}  (seed={a.seed}; „gorący” dział ukryty — znajdźcie go w E1)")


if __name__ == "__main__":
    main()
